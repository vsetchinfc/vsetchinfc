"""
doc1_lesson_plan.py
Generates: output/01_план_урока.docx
Audience: Teacher
Format: A4 portrait, 12pt
"""

import os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT_DIR = "output"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "01_план_урока.docx")

COLOR_NOUN = RGBColor(0x1E, 0x90, 0xFF)       # Blue
COLOR_ADJ = RGBColor(0x32, 0xCD, 0x32)        # Green
COLOR_VERB = RGBColor(0xFF, 0x45, 0x00)       # Red
COLOR_HEADER_BG = "1E90FF"
COLOR_LIGHT_BLUE = "DDEEFF"
COLOR_LIGHT_GREEN = "DDFFDD"
COLOR_LIGHT_RED = "FFE8E0"
COLOR_LIGHT_GREY = "F5F5F5"


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, border_color="000000", border_size="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), border_size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), border_color)
        tcBorders.append(border)
    tcPr.append(tcBorders)


def add_heading(doc, text, size=16, bold=True, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def add_paragraph_text(doc, text, size=12, bold=False, italic=False, color=None,
                        align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def set_table_border(table, border_color="AAAAAA"):
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), border_color)
        tblBorders.append(border)
    tblPr.append(tblBorders)


def create_lesson_plan():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc = Document()

    # --- Page setup: A4 portrait ---
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # =========================================================
    # HEADER
    # =========================================================
    add_heading(doc, "ПЛАН УРОКА", size=22, bold=True)
    add_paragraph_text(
        doc,
        "Тема: Части речи (Существительное, Прилагательное, Глагол)",
        size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER
    )
    add_paragraph_text(
        doc,
        "Класс: 2-й класс  |  Количество учеников: 10",
        size=12, align=WD_ALIGN_PARAGRAPH.CENTER
    )
    add_paragraph_text(
        doc,
        "Воскресная Русская Школа  |  Мельбурн, Австралия",
        size=12, align=WD_ALIGN_PARAGRAPH.CENTER
    )
    add_paragraph_text(
        doc,
        "Дата: _______________    Учитель: _______________",
        size=12, align=WD_ALIGN_PARAGRAPH.CENTER
    )
    add_paragraph_text(
        doc,
        "Продолжительность: 90 минут  |  Первое знакомство с темой",
        size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER
    )

    doc.add_paragraph()

    # =========================================================
    # LESSON GOALS — three-column table
    # =========================================================
    add_heading(doc, "ЦЕЛИ УРОКА", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    goals_table = doc.add_table(rows=2, cols=3)
    goals_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(goals_table, "888888")

    headers = ["ЗНАТЬ", "УМЕТЬ", "ХОТЕТЬ"]
    header_colors = [COLOR_LIGHT_BLUE, COLOR_LIGHT_GREEN, COLOR_LIGHT_RED]
    goals = [
        (
            "Существительное отвечает на КТО? ЧТО?\n"
            "Прилагательное — КАКОЙ? КАКАЯ? КАКОЕ?\n"
            "Глагол — ЧТО ДЕЛАЕТ?"
        ),
        "Находить и подчёркивать части речи в предложении нужной линией",
        "Применять знания самостоятельно дома",
    ]

    hdr_row = goals_table.rows[0]
    for i, (hdr, bg) in enumerate(zip(headers, header_colors)):
        cell = hdr_row.cells[i]
        set_cell_bg(cell, bg.replace("#", ""))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(hdr)
        run.bold = True
        run.font.size = Pt(12)

    data_row = goals_table.rows[1]
    for i, (goal, bg) in enumerate(zip(goals, header_colors)):
        cell = data_row.cells[i]
        set_cell_bg(cell, bg.replace("#", ""))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(goal)
        run.font.size = Pt(11)

    doc.add_paragraph()

    # =========================================================
    # TIMING TABLE
    # =========================================================
    add_heading(doc, "ТАЙМИНГ УРОКА", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    timing_data = [
        ("00–10 мин", "Организационный момент",
         "Приветствие. Игра «Назови предмет в классе» — дети называют всё, что видят. Учитель записывает на доску."),
        ("10–25 мин", "Существительное",
         "Показ Плаката 1. Скрипт: «Ребята, всё, что можно потрогать или назвать — это существительное. "
         "Спросим: КТО ЭТО? или ЧТО ЭТО?» Дети подходят к доске и подчёркивают слова одной линией."),
        ("25–40 мин", "Прилагательное",
         "Показ Плаката 2. Скрипт: «Прилагательное описывает предмет. Спросим: КАКОЙ? КАКАЯ? КАКОЕ? "
         "Оно как будто прилипает к существительному!» Игра «Опиши предмет тремя словами»."),
        ("40–55 мин", "Глагол",
         "Показ Плаката 3. Скрипт: «Глагол — это действие. Что делает кот? МЯУКАЕТ! Что делает мяч? ЛЕТИТ!» "
         "Игра «Покажи действие» — дети изображают глаголы жестами."),
        ("55–75 мин", "Проверочные задания",
         "Раздать Документ 3. Дети выполняют задания уровней 1–4 самостоятельно, уровень 5 — по желанию."),
        ("75–85 мин", "Проверка и разбор",
         "Проверить вместе. Разобрать типичные ошибки. Похвалить за старание."),
        ("85–90 мин", "Домашнее задание",
         "Раздать Документ 4. Объяснить задания 1–3. Задание 4 — по желанию (★)."),
    ]

    timing_table = doc.add_table(rows=len(timing_data) + 1, cols=3)
    timing_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(timing_table)

    # Header row
    col_headers = ["Время", "Этап", "Деятельность учителя и детей"]
    col_widths = [Cm(2.5), Cm(4), Cm(12)]
    timing_table.columns[0].width = col_widths[0]
    timing_table.columns[1].width = col_widths[1]
    timing_table.columns[2].width = col_widths[2]

    hdr_row = timing_table.rows[0]
    for i, hdr in enumerate(col_headers):
        cell = hdr_row.cells[i]
        set_cell_bg(cell, "2F4F8F")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(hdr)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    stage_colors = {
        "Существительное": "DDEEFF",
        "Прилагательное": "DDFFDD",
        "Глагол": "FFE8E0",
    }

    for row_idx, (time_val, stage, activity) in enumerate(timing_data, start=1):
        row = timing_table.rows[row_idx]
        bg = stage_colors.get(stage, "FFFFFF")

        cells_data = [time_val, stage, activity]
        for col_idx, val in enumerate(cells_data):
            cell = row.cells[col_idx]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.size = Pt(10)
            if col_idx == 1:
                run.bold = True

    doc.add_paragraph()

    # =========================================================
    # MATERIALS
    # =========================================================
    add_heading(doc, "МАТЕРИАЛЫ К УРОКУ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    materials = [
        "Распечатанный план урока (1 экз. для учителя)",
        "Плакат «Существительное» (Документ 2, лист 1) — 1 экз. формат A4 или A3",
        "Плакат «Прилагательное» (Документ 2, лист 2) — 1 экз.",
        "Плакат «Глагол» (Документ 2, лист 3) — 1 экз.",
        "Сводная таблица-шпаргалка (Документ 2, лист 4) — 10 экз. (по одному на ребёнка)",
        "Проверочные задания (Документ 3) — 10 экз.",
        "Домашняя работа (Документ 4) — 10 экз.",
        "Цветные карандаши или фломастеры для каждого ребёнка",
        "Линейка для подчёркивания",
    ]
    for item in materials:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.size = Pt(11)

    doc.add_paragraph()

    # =========================================================
    # METHODOLOGY NOTES
    # =========================================================
    add_heading(doc, "МЕТОДИЧЕСКИЕ ЗАМЕТКИ — СКРИПТЫ ДЛЯ УЧИТЕЛЯ", size=14, bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT)

    notes = [
        (
            "Для объяснения СУЩЕСТВИТЕЛЬНОГО:",
            COLOR_NOUN,
            "«Существительное — это слово, которое обозначает ПРЕДМЕТ или ЖИВОЕ СУЩЕСТВО. "
            "Ты можешь спросить: КТО ЭТО? или ЧТО ЭТО? Например: КТО ЭТО? — КОТ. ЧТО ЭТО? — МЯЧ. "
            "Это существительные! Мы подчёркиваем их ОДНОЙ прямой линией.»",
        ),
        (
            "Для объяснения ПРИЛАГАТЕЛЬНОГО:",
            COLOR_ADJ,
            "«Прилагательное описывает существительное — говорит нам, КАКОЙ предмет. "
            "Кот — КАКОЙ? РЫЖИЙ, ПУШИСТЫЙ, ВЕСЁЛЫЙ. Мяч — КАКОЙ? КРАСНЫЙ, КРУГЛЫЙ, БОЛЬШОЙ. "
            "Мы подчёркиваем прилагательные ВОЛНИСТОЙ линией, потому что они такие... красивые и разнообразные!»",
        ),
        (
            "Для объяснения ГЛАГОЛА:",
            COLOR_VERB,
            "«Глагол — это ДЕЙСТВИЕ. Что делает кот? БЕЖИТ, ПРЫГАЕТ, МЯУКАЕТ. "
            "Что делает солнце? СВЕТИТ. Глагол всегда можно изобразить — попробуйте! "
            "Мы подчёркиваем глаголы ДВОЙНОЙ линией — потому что они сильные и энергичные!»",
        ),
    ]

    for label, color, script in notes:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = color

        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.8)
        p2.paragraph_format.space_after = Pt(6)
        run2 = p2.add_run(script)
        run2.font.size = Pt(11)
        run2.italic = True

    doc.add_paragraph()

    # =========================================================
    # DIFFERENTIATION
    # =========================================================
    add_heading(doc, "ДИФФЕРЕНЦИАЦИЯ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    diff_items = [
        "Для детей со слабым русским: разрешить пользоваться шпаргалкой "
        "(Документ 2, лист 4) во время проверочного задания",
        "Для сильных детей: задание уровня 5 (★) в Документе 3 и задание 4 (★) в домашней работе",
    ]
    for item in diff_items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.size = Pt(11)

    # =========================================================
    # MARKING SYSTEM — reminder box
    # =========================================================
    doc.add_paragraph()
    add_heading(doc, "СИСТЕМА МАРКИРОВКИ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    mark_table = doc.add_table(rows=4, cols=3)
    mark_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(mark_table)

    mark_headers = ["Часть речи", "Вопросы", "Подчёркивание"]
    mark_hdr_row = mark_table.rows[0]
    for i, hdr in enumerate(mark_headers):
        cell = mark_hdr_row.cells[i]
        set_cell_bg(cell, "2F4F8F")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(hdr)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    mark_data = [
        ("СУЩЕСТВИТЕЛЬНОЕ", "КТО? ЧТО?", "Одна прямая линия  _____", "DDEEFF"),
        ("ПРИЛАГАТЕЛЬНОЕ", "КАКОЙ? КАКАЯ? КАКОЕ?", "Волнистая линия  ~~~~~", "DDFFDD"),
        ("ГЛАГОЛ", "ЧТО ДЕЛАЕТ?", "Двойная линия  \u2550\u2550\u2550\u2550\u2550", "FFE8E0"),
    ]

    for row_idx, (pos, questions, marking, bg) in enumerate(mark_data, start=1):
        row = mark_table.rows[row_idx]
        for col_idx, val in enumerate([pos, questions, marking]):
            cell = row.cells[col_idx]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(11)
            if col_idx == 0:
                run.bold = True

    doc.save(OUTPUT_FILE)
    print(f"✅  Создан: {OUTPUT_FILE}")


if __name__ == "__main__":
    create_lesson_plan()
